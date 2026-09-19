"""只读取原始工作簿；每条记录保留文件、工作表和行号。"""
import hashlib
import io
import re
from pathlib import Path
from zipfile import ZipFile, BadZipFile
import openpyxl
from scoring import CATEGORIES, clean, number

def text(v):
    if v is None: return ''
    if isinstance(v,float) and v.is_integer(): return str(int(v))
    return str(v).strip()

def parse_workbook(data, filename, mappings=None):
    if len(data)>25*1024*1024: raise ValueError('单个文件不得超过 25 MB')
    try:
        with ZipFile(io.BytesIO(data)) as z:
            if sum(x.file_size for x in z.infolist())>150*1024*1024: raise ValueError('工作簿解压后过大')
    except BadZipFile: raise ValueError('文件不是有效的 .xlsx 工作簿')
    w=openpyxl.load_workbook(io.BytesIO(data),data_only=True,read_only=True)
    digest=hashlib.sha256(data).hexdigest(); records=[]; sheets=[]; warnings=[]; mappings=mappings or {}
    for sheet_index,s in enumerate(w):
        sheet_key=digest+':'+s.title
        if mappings.get(sheet_key)=='__SKIP__':
            warnings.append(f'{s.title}：用户跳过'); continue
        if any(k in s.title for k in ['填写说明','示例','说明页']):
            warnings.append(f'{s.title}：跳过说明/示例页'); continue
        if s.max_row>50000 or s.max_column>100: raise ValueError('工作表范围过大，请移除多余空白行列')
        header=None; category=''; event=''; date=''; count=0; metadata=[]; detected=set(); matched_by=''; category_col_idx=None
        for rownum,row in enumerate(s.iter_rows(values_only=True),1):
            vals=[text(v) for v in row]; normalized=[clean(v) for v in vals]
            if '学号' in normalized and '姓名' in normalized:
                header={v:i for i,v in enumerate(normalized) if v}
                # 检测汇总式表格的分数类别列
                for name in ['类别','活动性质','分数类别','性质','附加分类别','计分类别']:
                    if name in header:
                        category_col_idx=header[name]
                        warnings.append(f'{s.title}：检测到汇总式表格（类别列：{name}），将按列值自动拆分')
                        break
                mapped=mappings.get(digest+':'+s.title)
                if mapped is not None and mapped not in [*CATEGORIES,'__ROWS__']: raise ValueError('手动指定的类别无效')
                if mapped and mapped!='__ROWS__': category=mapped; matched_by='手动指定'
                elif len(detected)==1: category=next(iter(detected)); matched_by='表内活动性质'
                elif len(detected)>1: category=''; matched_by='表内勾选了多个类别'
                else:
                    title=clean(s.title)
                    aliases={'专业学术':['专业学术','学术类'],'文体科技实践活动':['文体科技实践','文体活动','文体类'],'志愿服务':['志愿服务','志愿时长'],'社会实践':['社会实践'],'学生干部':['学生干部','干部任职'],'荣誉表彰':['荣誉表彰','各级荣誉','荣誉类'],'其他':['其他'],'专业论文':['专业论文'],'非专业作品':['非专业作品'],'无偿献血':['献血'],'处罚分':['处罚分','惩罚分']}
                    matches=[cat for cat,terms in aliases.items() if any(term in title for term in terms)]
                    if len(matches)==1: category=matches[0]; matched_by='工作表名称'
                continue
            if header is None:
                metadata.extend(vals)
                for val in vals:
                    if '项目名称' in val and re.search('[：:]',val): event=re.split('[：:]',val,1)[1].strip()
                    if '项目时间' in val and re.search('[：:]',val): date=re.split('[：:]',val,1)[1].strip()
                    for cat in CATEGORIES:
                        if re.search(re.escape(cat)+r'\s*[\[【（(]\s*[√✓✔]',val): detected.add(cat)
                    if re.search(r'文体活动\s*\[\s*[√✓✔]',val): detected.add('文体科技实践活动')
                continue
            def get(*names):
                for name in names:
                    if name in header and header[name]<len(vals): return vals[header[name]]
                return ''
            sid=get('学号'); name=get('姓名'); college=get('院系','学院','学部(院、系)'); klass=get('班级')
            award=get('奖项/服务时长/职务','奖项','服务时长','职务'); score=get('分值','分数'); note=get('备注')
            if not any([sid,name,award,score,klass]): continue
            # 尾部注释不进入人员记录；有班级、奖项但无学号的集体条目则保留待审核。
            if not any([sid,name,award]) or (name in ['姓名','签字']): continue
            # 汇总式表格：从类别列读取当前行的类别
            if category_col_idx is not None and not (mappings.get(sheet_key) in CATEGORIES):
                row_category=clean(vals[category_col_idx]) if category_col_idx<len(vals) else ''
                aliases={'专业学术':['专业学术','学术类'],'文体科技实践活动':['文体科技实践','文体活动','文体类'],'志愿服务':['志愿服务','志愿时长'],'社会实践':['社会实践'],'学生干部':['学生干部','干部任职'],'荣誉表彰':['荣誉表彰','各级荣誉','荣誉类'],'其他':['其他'],'专业论文':['专业论文'],'非专业作品':['非专业作品'],'无偿献血':['献血'],'处罚分':['处罚分','惩罚分']}
                matched=[cat for cat,terms in aliases.items() if any(term in row_category for term in terms)]
                cat=matched[0] if len(matched)==1 else (row_category if row_category in CATEGORIES else '')
            else:
                cat=category or '其他'
            issue=''
            if score and number(score) is None: issue='原分值非数字或公式缺少缓存结果'
            if not score: issue='原表未填写分值（可能为未缓存公式）'
            if not cat: issue=(issue+'；' if issue else '')+'未识别活动性质'
            extra='；'.join(v for j,v in enumerate(vals) if j>=10 and v)
            if extra: note+='；'+extra
            records.append({'id':digest[:14]+'-'+str(sheet_index)+'-'+str(rownum),'source':Path(filename).name,'sheet':s.title,'row':rownum,'student_id':sid,'name':name,'college':college,'class_name':klass,'category':cat or '其他','award':award,'level':get('级别'),'kind':get('个人/集体'),'original_score':number(score),'raw_score':score,'event':event or note,'event_group':'','event_date':date,'source_note':note,'review':'auto','review_note':'','import_issue':issue})
            count+=1
        sheets.append({'name':s.title,'count':count,'category':category,'matched_by':('按每行类别匹配；未知项进入待审核' if category_col_idx is not None else matched_by),'mixed':category_col_idx is not None,'needs_mapping':bool(count and not category and category_col_idx is None),'key':digest+':'+s.title})
        if header is None: warnings.append(f'{s.title}：未找到学号/姓名表头')
    w.close()
    return {'hash':digest,'name':Path(filename).name,'sheets':sheets,'count':len(records),'warnings':warnings},records
