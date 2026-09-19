"""只提取人员信息；不导入分值，不覆盖已有名册。"""
import base64
import hashlib
import io
import json
import re
from pathlib import Path
from zipfile import ZipFile, BadZipFile
import openpyxl
from importer import text
import registry

ALIASES={'student_id':['学号','学生学号'], 'name':['姓名','学生','学生姓名'], 'college':['院系','学院','院系所','院系名称','所属院系','学部(院、系)','学部（院、系）'], 'class_name':['班级','班级名称','行政班','行政班级'], 'code':['院系代码','院系所代码','学院代码']}

def normalize(v):return re.sub(r'\s+','',text(v)).rstrip('：:')

def read_workbook(raw,filename):
    if len(raw)>25*1024*1024:raise ValueError('单文件最多 25 MB')
    try:
        with ZipFile(io.BytesIO(raw)) as z:
            if sum(i.file_size for i in z.infolist())>150*1024*1024:raise ValueError('工作簿解压后过大')
    except BadZipFile:raise ValueError('不是有效的 .xlsx 工作簿')
    rows=[];warnings=[]
    w=openpyxl.load_workbook(io.BytesIO(raw),data_only=True,read_only=True)
    try:
        for sheet in w:
            if any(x in sheet.title for x in ['填写说明','示例','说明页']):continue
            if (sheet.max_row or 0)>50000 or (sheet.max_column or 0)>100:raise ValueError('工作表范围过大')
            header=None
            for num,values in enumerate(sheet.iter_rows(values_only=True),1):
                values=list(values);norm=[normalize(v) for v in values]
                found={k:next((i for i,v in enumerate(norm) if v in names),None) for k,names in ALIASES.items()}
                if found['student_id'] is not None and found['name'] is not None:
                    header=found;continue
                if header is None:continue
                item={k:text(values[i]) if i is not None and i<len(values) else '' for k,i in header.items()}
                if not item['student_id'] and not item['name']:continue
                item.update(source=Path(filename).name,sheet=sheet.title,row=num);rows.append(item)
            if header is None:warnings.append(f'{filename} / {sheet.title}：未找到学号和姓名表头，已跳过')
    finally:w.close()
    return rows,warnings

def preview(c,body,source_rows=None):
    rows=[];warnings=[]
    if body.get('from_existing'):
        rows=source_rows if source_rows is not None else [json.loads(r[0]) for r in c.execute('SELECT current FROM records ORDER BY id')]
    else:
        files=body.get('files',[])
        if not 1<=len(files)<=20:raise ValueError('请选择 1～20 个 .xlsx 工作簿')
        for f in files:
            if not f['name'].lower().endswith('.xlsx'):raise ValueError('当前支持 .xlsx 工作簿')
            found,notes=read_workbook(base64.b64decode(f['data'],validate=True),f['name']);rows+=found;warnings+=notes
    master=registry.get(c);bycode={d['code']:d for d in master['departments'] if d['code']};byname={d['name']:d for d in master['departments']};existing={s['student_id']:s for s in master['students']}
    recycled={r['student_id'] for r in c.execute('SELECT student_id FROM roster WHERE deleted=1')}
    recycled_classes={(r['college'],r['name']) for r in registry.get_recycle(c)['classes']}
    groups={}
    for row in rows:
        r={k:text(row.get(k,'')) for k in ['student_id','name','college','class_name','source','sheet','row','code']}
        dep=bycode.get(r['code'] or r['college'])
        if dep:r['college']=dep['name']
        r['class_name']=registry.canonical_class(c,r['college'],r['class_name'])
        scope=body.get('college_filter')
        if scope and scope not in ['*','全部学院'] and r['college'] and r['college']!=scope:continue
        groups.setdefault(r['student_id'] or 'missing:'+str(len(groups)),[]).append(r)
    items=[]
    for key,group in groups.items():
        r=group[0];item={**r,'sources':[{'source':x['source'],'sheet':x['sheet'],'row':x['row']} for x in group],'variants':[], 'status':'ready','reason':'新增学生；相同信息已合并' if len(group)>1 else '新增学生'}
        variants=list(dict.fromkeys((x['name'],x['college'],x['class_name']) for x in group));item['variants']=[dict(zip(['name','college','class_name'],x)) for x in variants]
        if not re.fullmatch(r'\d{10,12}',r['student_id']) or any(not x[k] for x in group for k in ['name','college','class_name']):item.update(status='invalid',reason='学号须为 10～12 位数字，且姓名、院系、班级完整')
        elif len(variants)>1:item.update(status='conflict',reason='同一学号的姓名、院系或班级不一致，请先核对')
        elif r['student_id'] in recycled or (r['college'],r['class_name']) in recycled_classes:item.update(status='conflict',reason='学生或班级在回收站，请先恢复或核对')
        elif r['student_id'] in existing:
            old=existing[r['student_id']];same=all(old[k]==r[k] for k in ['name','college','class_name'])
            item.update(status='existing' if same else 'conflict',reason='名册已存在，跳过（不改变停用状态）' if same else '与已有名册不一致，不覆盖',existing=old)
        elif r['college'] not in byname:item.update(status='invalid',reason='院系未匹配固定字典，请更正表格；确需新增时修改 departments.json')
        elif not byname[r['college']]['active']:item.update(status='invalid',reason='院系已停用，请先核对')
        elif any(x['college']==r['college'] and x['name']==r['class_name'] and not x['active'] for x in master['classes']):item.update(status='invalid',reason='对应班级已停用，请先核对')
        if body.get('from_existing') and re.match(r'^20\d{2}2\d{7}$',r['student_id']):item.update(status='invalid',reason='申报学号疑似研究生，请核对后使用独立名册导入')
        item['key']=key;items.append(item)
    result={'items':items,'warnings':warnings,'rows':len(rows),'ready':sum(x['status']=='ready' for x in items)}
    result['fingerprint']=hashlib.sha256(json.dumps(result,ensure_ascii=False,sort_keys=True).encode()).hexdigest()
    return result

def commit(c,body,log,actor,source_rows=None):
    report=preview(c,body,source_rows)
    if body.get('fingerprint')!=report['fingerprint']:raise ValueError('名册或申报已变化，请重新预览后导入')
    selected=set(body.get('selected',[]));ready={i['key']:i for i in report['items'] if i['status']=='ready'}
    if not selected or not selected<=ready.keys():raise ValueError('请选择预览中可新增的学生')
    for key in sorted(selected):
        r=ready[key];dep=c.execute('SELECT id FROM departments WHERE name=?',(r['college'],)).fetchone()[0]
        c.execute('INSERT OR IGNORE INTO classes(department_id,name) VALUES(?,?)',(dep,r['class_name']))
        cls=c.execute('SELECT id FROM classes WHERE department_id=? AND name=?',(dep,r['class_name'])).fetchone()[0]
        ident,before,after=registry.save(c,{'kind':'student','class_id':cls,'student_id':r['student_id'],'name':r['name']})
        log(c,actor,'导入学生名册',ident,before,{'student':after,'sources':r['sources']})
    return {'added':len(selected)}
